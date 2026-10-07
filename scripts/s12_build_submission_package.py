"""v3 -> IJC submission package.

Collects the submission files into ../15_Submission_Package_IJC_20260912 with the names the
Wiley Authors portal asks for, converts the three markdown sources to .docx with the shared
converter, assembles the ten supplementary figures with their legends into one PDF, merges the
thirteen supplementary tables into one workbook, and writes a manifest with SHA-256 sums.
Run after s05-s09 have all passed."""
import hashlib
import os
import re
import shutil
import sys

import openpyxl
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer
from PIL import Image as PILImage

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from md2docx_v3 import md_to_docx

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
V3 = os.path.join(BASE, "14_Revision_v3_20260806")
PKG = os.path.join(BASE, "15_Submission_Package_IJC_20260912")
MS_MD = os.path.join(V3, "Manuscript_IJC_v3_DLL3_Immune_TME_20260806.md")
CL_MD = os.path.join(V3, "Cover_Letter_IJC_v3_20260806.md")
LEG_MD = os.path.join(V3, "Figure_Legends_IJC_v3_20260806.md")

for sub in ["04_Figures", "05_Tables", "06_Supplementary"]:
    os.makedirs(os.path.join(PKG, sub), exist_ok=True)

# 1. text files ------------------------------------------------------------------------------
md_to_docx(MS_MD, os.path.join(PKG, "01_Main_Document_Manuscript.docx"))
md_to_docx(CL_MD, os.path.join(PKG, "02_Cover_Letter.docx"))
md_to_docx(LEG_MD, os.path.join(PKG, "03_Figure_Legends.docx"))

# 2. main figures (line art: vector PDF at 600 dpi raster fallback, plus TIFF) and tables ----
FIGS = {1: "Figure1_v3_Study_Design_and_Cohort_Flow",
        2: "Figure2_v3_Antigen_Presentation_and_Leukocyte_Content",
        3: "Figure3_v3_TNFRSF9_CD8_Phenotype",
        4: "Figure4_v3_Epithelial_APM_and_TNFRSF9_CD8",
        5: "Figure5_v3_What_Is_And_Is_Not_Established"}
for n, stem in FIGS.items():
    for ext in ("pdf", "tiff"):
        shutil.copy2(os.path.join(V3, "figures", f"{stem}.{ext}"), os.path.join(PKG, "04_Figures", f"Figure{n}.{ext}"))
shutil.copy2(os.path.join(V3, "figures", "Graphical_Abstract_v3.pdf"),
             os.path.join(PKG, "07_Graphical_Abstract.pdf"))
for src, dst in [("Table1_v3_Questions_and_Dataset_Roles.docx", "Table1.docx")]:
    shutil.copy2(os.path.join(V3, "tables", src), os.path.join(PKG, "05_Tables", dst))

# 3. supplementary figures S1-S10 with legends, one PDF ---------------------------------------
import matplotlib
ttf = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")
for name, fn in [("DejaVu", "DejaVuSans.ttf"), ("DejaVu-Bold", "DejaVuSans-Bold.ttf"),
                 ("DejaVu-Italic", "DejaVuSans-Oblique.ttf"), ("DejaVu-BoldItalic", "DejaVuSans-BoldOblique.ttf")]:
    pdfmetrics.registerFont(TTFont(name, os.path.join(ttf, fn)))
pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu-Italic", boldItalic="DejaVu-BoldItalic")
body = ParagraphStyle("b", fontName="DejaVu", fontSize=8.5, leading=11.5)
head = ParagraphStyle("h", fontName="DejaVu-Bold", fontSize=10, leading=13, spaceAfter=4)


def md_inline(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"\*(.+?)\*", r"<i>\1</i>", t)
    return t.replace("`", "")


leg = open(LEG_MD, encoding="utf-8").read().split("## Supplementary Figures", 1)[1]
legends = {int(m.group(1)): m.group(0).strip()
           for m in re.finditer(r"\*\*Supplementary Figure S(\d+)\..*?(?=\n\*\*Supplementary Figure S|\Z)", leg, flags=re.S)}
assert sorted(legends) == list(range(1, 11)), sorted(legends)
W, H = A4
story = [Paragraph("Supplementary Figures S1–S10", head),
         Paragraph(md_inline("DLL3 expression and its associations with antigen-presentation genes and *TNFRSF9*-expressing CD8 T cells in small cell lung cancer. Kobayashi et al."), body),
         PageBreak()]
figdir = os.path.join(V3, "figures")
for n in range(1, 11):
    png = next(f for f in os.listdir(figdir) if f.startswith(f"SuppFigureS{n}_") and f.endswith(".png"))
    with PILImage.open(os.path.join(figdir, png)) as im:
        w, h = im.size
    maxw, maxh = W - 30 * mm, H - 95 * mm
    scale = min(maxw / w, maxh / h)
    story += [Paragraph(f"Supplementary Figure S{n}", head),
              Image(os.path.join(figdir, png), width=w * scale, height=h * scale),
              Spacer(1, 4 * mm), Paragraph(md_inline(legends[n]), body), PageBreak()]
SimpleDocTemplate(os.path.join(PKG, "06_Supplementary", "Supplementary_Figures_S1-S10.pdf"), pagesize=A4,
                  leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                  title="Supplementary Figures S1-S10", author="Kobayashi N et al.").build(story)

# 4. supplementary tables S1-S13, one workbook (sheet per table / sub-table) ------------------
out = openpyxl.Workbook(); out.remove(out.active)
supp = os.path.join(V3, "supplementary")
for fn in sorted(os.listdir(supp), key=lambda s: int(re.search(r"_S(\d+)_", s).group(1))):
    if not fn.endswith(".xlsx"):
        continue
    src = openpyxl.load_workbook(os.path.join(supp, fn))
    for ws in src.worksheets:
        dst = out.create_sheet(ws.title[:31])
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d = dst.cell(row=c.row, column=c.column, value=c.value)
                    d.font = c.font.copy(); d.alignment = c.alignment.copy(); d.number_format = c.number_format
        for col, dim in ws.column_dimensions.items():
            if dim.width:
                dst.column_dimensions[col].width = dim.width
        dst.freeze_panes = ws.freeze_panes
out.save(os.path.join(PKG, "06_Supplementary", "Supplementary_Tables_S1-S13.xlsx"))
for fn in os.listdir(supp):
    if fn.endswith(".xlsx"):
        shutil.copy2(os.path.join(supp, fn), os.path.join(PKG, "06_Supplementary", "individual_tables", fn)
                     if os.path.isdir(os.path.join(PKG, "06_Supplementary", "individual_tables"))
                     else (os.makedirs(os.path.join(PKG, "06_Supplementary", "individual_tables"), exist_ok=True)
                           or os.path.join(PKG, "06_Supplementary", "individual_tables", fn)))

# 5. manifest ---------------------------------------------------------------------------------
lines = ["# Submission package manifest", "", f"Built by s12_build_submission_package.py from {os.path.basename(V3)}.", "",
         "| File | Size (KB) | SHA-256 (first 16) |", "|---|---:|---|"]
for root, _, files in sorted(os.walk(PKG)):
    for f in sorted(files):
        if f.startswith("00_") or f == ".DS_Store":
            continue
        p = os.path.join(root, f)
        h = hashlib.sha256(open(p, "rb").read()).hexdigest()[:16]
        lines.append(f"| {os.path.relpath(p, PKG)} | {os.path.getsize(p) / 1024:,.0f} | {h} |")
open(os.path.join(PKG, "00_MANIFEST.md"), "w", encoding="utf-8").write("\n".join(lines) + "\n")
print("\n".join(lines))
