"""Markdown -> .docx for the IJC submission files (manuscript, cover letter, figure legends).

Shared by s08 and s12. Handles the markdown actually used in these files: '# ' / '## '
headings, **bold**, *italic*, italics nested inside bold, `code` spans (rendered as plain
text: they mark sample identifiers), escaped \\* and the '---' rule. Anything else is
written verbatim."""
import re
from docx import Document
from docx.shared import Pt

_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|\\\*)")


def _runs(p, text, bold=False, italic=False):
    for tok in _TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            _runs(p, tok[2:-2], bold=True, italic=italic)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            _runs(p, tok[1:-1], bold=bold, italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1]); r.bold, r.italic = bold, italic
        elif tok == "\\*":
            r = p.add_run("*"); r.bold, r.italic = bold, italic
        else:
            r = p.add_run(tok); r.bold, r.italic = bold, italic


def md_to_docx(md_path, docx_path, font="Times New Roman", size=11):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    d = Document()
    d.styles["Normal"].font.name = font
    d.styles["Normal"].font.size = Pt(size)
    for ln in lines:
        s = ln.rstrip()
        if not s.strip() or s.strip() == "---":
            continue
        if s.startswith("# "):
            p = d.add_paragraph(); _runs(p, s[2:].strip(), bold=True)
            for r in p.runs: r.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(14)
        elif s.startswith("## "):
            p = d.add_paragraph(); _runs(p, s[3:].strip(), bold=True)
            for r in p.runs: r.font.size = Pt(12)
            p.paragraph_format.space_before = Pt(10)
        else:
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            _runs(p, s)
    d.save(docx_path)
    return docx_path
